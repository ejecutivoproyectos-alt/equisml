import os
import json
import re
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from xml.etree import ElementTree as ET

import streamlit as st
from dotenv import load_dotenv
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openai import OpenAI

from app.db.database import SessionLocal
from app.models.empresa import Empresa
from app.models.empresa_estilo_word import EmpresaEstiloWord
from app.models.empresa_plantilla_asignacion import EmpresaPlantillaAsignacion
from app.models.empresa_plantilla_word import EmpresaPlantillaWord
from utils.calendario_mexico import calcular_fecha_larga_habil_mexico


load_dotenv()

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

COLOR_PRINCIPAL_FALLBACK = (31, 58, 95)
COLOR_TEXTO_FALLBACK = (0, 0, 0)
FUENTE_BASE_FALLBACK = "Calibri"
TAMANIO_BASE_FALLBACK = 11


MESES_ES = {
    1: "enero",
    2: "febrero",
    3: "marzo",
    4: "abril",
    5: "mayo",
    6: "junio",
    7: "julio",
    8: "agosto",
    9: "septiembre",
    10: "octubre",
    11: "noviembre",
    12: "diciembre",
}

MESES_SELECT = [
    (1, "Enero"),
    (2, "Febrero"),
    (3, "Marzo"),
    (4, "Abril"),
    (5, "Mayo"),
    (6, "Junio"),
    (7, "Julio"),
    (8, "Agosto"),
    (9, "Septiembre"),
    (10, "Octubre"),
    (11, "Noviembre"),
    (12, "Diciembre"),
]


def extraer_fecha_xml_cfdi(archivo_xml):
    try:
        archivo_xml.seek(0)
        root = ET.parse(archivo_xml).getroot()
    except ET.ParseError:
        raise ValueError(f"El archivo {archivo_xml.name} no es un XML valido.")

    fecha_texto = root.attrib.get("Fecha")

    if not fecha_texto:
        for elemento in root.iter():
            fecha_texto = elemento.attrib.get("Fecha")
            if fecha_texto:
                break

    if not fecha_texto:
        raise ValueError(f"No se encontro la etiqueta Fecha en {archivo_xml.name}.")

    fecha_sin_hora = fecha_texto.split("T", 1)[0]

    try:
        return datetime.strptime(fecha_sin_hora, "%Y-%m-%d").date()
    except ValueError:
        raise ValueError(
            f"La fecha del archivo {archivo_xml.name} no tiene formato valido: {fecha_texto}"
        )


def obtener_fechas_ordenadas_xml(archivos_xml):
    fechas = [extraer_fecha_xml_cfdi(archivo) for archivo in archivos_xml]
    return sorted(fechas)


def calcular_fecha_documento_desde_xml(archivos_xml):
    fechas_ordenadas = obtener_fechas_ordenadas_xml(archivos_xml)

    if not fechas_ordenadas:
        raise ValueError("Sube al menos un XML para calcular la fecha del documento.")

    fecha_inicio = fechas_ordenadas[0]
    fecha_inicio_str = fecha_inicio.strftime("%d/%m/%Y")
    fecha_documento = calcular_fecha_larga_habil_mexico(
        fecha_str=fecha_inicio_str,
        dias_restar=4,
        dias_sumar=0,
    )

    return fecha_documento, fechas_ordenadas


def formatear_fecha_corta(fecha):
    return fecha.strftime("%d/%m/%Y")


def convertir_color_bd_a_rgb(color):
    if not color:
        return COLOR_PRINCIPAL_FALLBACK

    color = color.strip().lstrip("#")

    if len(color) != 6:
        return COLOR_PRINCIPAL_FALLBACK

    try:
        return tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return COLOR_PRINCIPAL_FALLBACK


def aplicar_fuente_run(run, fuente):
    run.font.name = fuente
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), fuente)
    r_fonts.set(qn("w:hAnsi"), fuente)
    r_fonts.set(qn("w:eastAsia"), fuente)
    r_fonts.set(qn("w:cs"), fuente)


def obtener_empresas_triple_a():
    db = SessionLocal()

    try:
        resultados = (
            db.query(Empresa, EmpresaPlantillaAsignacion, EmpresaPlantillaWord)
            .outerjoin(
                EmpresaPlantillaAsignacion,
                EmpresaPlantillaAsignacion.empresa_externa_id == Empresa.id,
            )
            .outerjoin(
                EmpresaPlantillaWord,
                EmpresaPlantillaWord.id == EmpresaPlantillaAsignacion.plantilla_id,
            )
            .filter(Empresa.tipo_empresa == "TRIPLE_AAA")
            .order_by(Empresa.nombre.asc())
            .all()
        )

        return [
            {
                "id": empresa.id,
                "nombre": empresa.nombre,
                "razon_social": empresa.razon_social or empresa.nombre,
                "color_primario": empresa.color_primario,
                "color_secundario": empresa.color_secundario,
                "color_acento": empresa.color_acento,
                "membrete_path": asignacion.membrete_path if asignacion else None,
                "tipografia_base": plantilla.tipografia_base if plantilla else FUENTE_BASE_FALLBACK,
                "tamanio_base": plantilla.tamanio_base if plantilla else TAMANIO_BASE_FALLBACK,
            }
            for empresa, asignacion, plantilla in resultados
        ]
    finally:
        db.close()


def obtener_estilos_word_empresa(empresa_id):
    db = SessionLocal()

    try:
        estilos = (
            db.query(EmpresaEstiloWord)
            .filter(EmpresaEstiloWord.empresa_id == empresa_id)
            .all()
        )

        return {
            estilo.clave_estilo: {
                "tipografia": estilo.tipografia,
                "tamanio_letra": estilo.tamanio_letra,
                "negrita": estilo.negrita,
                "cursiva": estilo.cursiva,
                "alineacion": estilo.alineacion,
            }
            for estilo in estilos
        }
    finally:
        db.close()


def obtener_cliente_openai():
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("No se encontro OPENAI_API_KEY en el archivo .env.")

    return OpenAI(api_key=api_key)


def llamar_openai_texto(prompt, modelo="gpt-4.1-mini"):
    client = obtener_cliente_openai()
    respuesta = client.responses.create(
        model=modelo,
        input=prompt,
    )
    return respuesta.output_text.strip()


def generar_metodologia_ia(nombre_servicio, nombre_cliente):
    prompt = f"""
Actua como consultor empresarial senior en Mexico.

Con base en el siguiente nombre de programa o servicio:
{nombre_servicio}

Genera una metodologia para un documento de seguimiento bimestral.

Instrucciones:
1. Define exactamente 1 tema principal breve, formal y especifico, relacionado directamente con el nombre del programa.
2. Genera exactamente 4 subtemas relacionados con ese tema.
3. El tema y los subtemas deben ser actividades o lineas de trabajo concretas, profesionales y utiles para seguimiento.
4. No uses explicaciones adicionales.

Entrega exclusivamente un JSON valido con esta estructura:
{{
  "tema": "Tema principal.",
  "subtemas": [
    "Subtema 1.",
    "Subtema 2.",
    "Subtema 3.",
    "Subtema 4."
  ]
}}
"""

    texto = llamar_openai_texto(prompt)
    texto_limpio = limpiar_respuesta_json_ia(texto)

    try:
        data = json.loads(texto_limpio)
        tema = str(data.get("tema", "")).strip()
        subtemas = [
            str(item).strip()
            for item in data.get("subtemas", [])
            if str(item).strip()
        ]
    except Exception:
        tema = nombre_servicio.upper()
        subtemas = [
            linea.strip(" -\t")
            for linea in texto.splitlines()
            if linea.strip(" -\t")
        ]

    if not tema:
        tema = nombre_servicio.upper()

    if len(subtemas) < 4:
        subtemas.extend([
            "Integracion y organizacion de informacion relevante.",
            "Seguimiento de avances y cumplimiento de actividades.",
            "Revision de observaciones y ajustes del servicio.",
            "Control administrativo de evidencias y resultados.",
        ])

    return {
        "tema": tema,
        "subtemas": subtemas[:4],
    }


def generar_introduccion_entregable_ia(nombre_servicio, tema_metodologia, nombre_cliente):
    prompt = f"""
Actua como consultor corporativo senior en Mexico.

Redacta una introduccion formal para un entregable documental de servicios.

Datos:
- Nombre del programa: {nombre_servicio}
- Empresa receptora del servicio: {nombre_cliente}
- Tema principal abordado: {tema_metodologia}

Instrucciones:
1. Explica de forma profesional de que tratara el documento.
2. Relaciona el contenido con el nombre del programa y el tema principal.
3. Menciona que el documento integra el seguimiento, desarrollo y evidencia de las actividades realizadas.
4. Usa un tono ejecutivo, claro y corporativo.
5. No inventes fechas, montos ni nombres adicionales.
6. No uses listas, viñetas ni encabezados.
7. Escribe entre 120 y 170 palabras.

Entrega exclusivamente el texto de la introduccion.
"""

    return llamar_openai_texto(prompt)


def generar_objetivos_entregable_ia(nombre_servicio, tema_metodologia, nombre_cliente):
    prompt = f"""
Actua como consultor corporativo senior en Mexico.

Genera los objetivos para un entregable documental de servicios.

Datos:
- Nombre del programa: {nombre_servicio}
- Empresa receptora del servicio: {nombre_cliente}
- Tema principal abordado: {tema_metodologia}

Instrucciones:
1. Genera un objetivo general en una sola oracion, con una estructura similar a:
"Al finalizar el servicio {nombre_cliente}, contara con las herramientas necesarias que le permita..."
2. Genera exactamente 3 objetivos especificos.
3. Cada objetivo especifico debe iniciar exactamente con:
"Al termino del servicio, {nombre_cliente} ..."
4. No escribas el nombre del programa despues de la frase "Al termino del servicio".
5. Relaciona todos los objetivos con el nombre del programa y el tema principal.
6. No uses listas, viñetas ni explicaciones fuera del JSON.

Entrega exclusivamente un JSON valido con esta estructura:
{{
  "objetivo_general": "Texto del objetivo general.",
  "objetivos_especificos": [
    "Objetivo especifico 1.",
    "Objetivo especifico 2.",
    "Objetivo especifico 3."
  ]
}}
"""

    texto = llamar_openai_texto(prompt)
    texto_limpio = limpiar_respuesta_json_ia(texto)

    try:
        data = json.loads(texto_limpio)
        objetivo_general = str(data.get("objetivo_general", "")).strip()
        objetivos_especificos = [
            str(item).strip()
            for item in data.get("objetivos_especificos", [])
            if str(item).strip()
        ]
    except Exception:
        objetivo_general = (
            f"Al finalizar el servicio {nombre_cliente.upper()} contara con las herramientas "
            "necesarias que le permitan fortalecer sus procesos, mejorar su eficiencia "
            "operativa y contribuir al desarrollo de su productividad."
        )
        objetivos_especificos = []

    if not objetivo_general:
        objetivo_general = (
            f"Al finalizar el servicio {nombre_cliente.upper()} contara con las herramientas "
            "necesarias que le permitan fortalecer sus procesos, mejorar su eficiencia "
            "operativa y contribuir al desarrollo de su productividad."
        )

    if len(objetivos_especificos) < 3:
        objetivos_especificos.extend([
            f"Al termino del servicio, {nombre_cliente.upper()} identificara las herramientas y procesos clave relacionados con {tema_metodologia.lower()}.",
            f"Al termino del servicio, {nombre_cliente.upper()} reconocera los beneficios operativos derivados de la correcta aplicacion del servicio.",
            f"Al termino del servicio, {nombre_cliente.upper()} podra fortalecer sus actividades internas mediante criterios de seguimiento, control y mejora continua.",
        ])

    objetivos_especificos = [
        normalizar_objetivo_especifico(objetivo, nombre_cliente)
        for objetivo in objetivos_especificos
    ]

    return {
        "objetivo_general": objetivo_general,
        "objetivos_especificos": objetivos_especificos[:3],
    }


def normalizar_objetivo_especifico(objetivo, nombre_cliente):
    objetivo = objetivo.strip()
    cliente = nombre_cliente.upper()
    patron = r"^Al t[ée]rmino del servicio(?:\s+[^,]+)?,\s*"
    objetivo = re.sub(
        patron,
        f"Al termino del servicio, {cliente} ",
        objetivo,
        flags=re.IGNORECASE,
    )
    objetivo = objetivo.replace(f"{cliente} {cliente}", cliente)

    if not objetivo.lower().startswith("al termino del servicio,"):
        objetivo = f"Al termino del servicio, {cliente} {objetivo}"

    return objetivo


def obtener_conceptos_metodologia(metodologia):
    conceptos = [metodologia["tema"]]
    conceptos.extend(metodologia["subtemas"])
    return [concepto for concepto in conceptos if concepto]


def generar_desarrollo_conceptos_entregable_ia(nombre_servicio, conceptos, nombre_cliente):
    desarrollos = []

    for concepto in conceptos:
        prompt = f"""
Actua como consultor corporativo senior en Mexico.

Redacta el desarrollo de un concepto para un entregable documental de servicios.

Datos:
- Nombre del programa: {nombre_servicio}
- Empresa receptora del servicio: {nombre_cliente}
- Concepto a desarrollar: {concepto}

Instrucciones:
1. Redacta un texto amplio, formal y profesional sobre el concepto indicado.
2. El texto debe explicar que se abordo, como se desarrollo, que actividades se consideran, que valor aporta a la empresa y como se relaciona con el programa.
3. Debe tener una extension aproximada equivalente a 2 paginas de Word, entre 850 y 1050 palabras.
4. Divide el texto en parrafos claros de 5 a 8 lineas.
5. No uses listas, viñetas, tablas ni encabezados internos.
6. No inventes fechas, montos, personas ni evidencias especificas.
7. No repitas literalmente el mismo cierre en cada parrafo.

Entrega exclusivamente el texto desarrollado.
"""

        texto = llamar_openai_texto(prompt)
        desarrollos.append({
            "concepto": concepto,
            "texto": texto,
        })

    return desarrollos


def generar_referencias_entregable_ia(nombre_servicio, desarrollo_conceptos):
    conceptos = "\n".join(
        f"- {desarrollo['concepto']}"
        for desarrollo in desarrollo_conceptos
    )

    prompt = f"""
Actua como documentalista academico especializado en formato APA 7.

Genera referencias bibliograficas para un entregable de servicios.

Nombre del programa:
{nombre_servicio}

Conceptos desarrollados:
{conceptos}

Instrucciones:
1. Genera entre 1 y 2 referencias por cada concepto.
2. Usa formato APA 7.
3. Da preferencia a paginas web, instituciones, normas, guias o articulos disponibles en idioma espanol.
4. No inventes autores, titulos, instituciones, URLs, fechas ni editoriales.
5. Si no tienes certeza de que una fuente existe, no la incluyas.
6. No agregues explicaciones, encabezados ni numeracion.
7. Entrega una referencia por linea.
"""

    texto = llamar_openai_texto(prompt)
    referencias = [
        linea.strip().lstrip("-•0123456789. ")
        for linea in texto.splitlines()
        if linea.strip()
    ]

    return referencias


def limpiar_respuesta_json_ia(texto):
    texto = texto.strip()

    if texto.startswith("```"):
        texto = texto.replace("```json", "", 1).replace("```JSON", "", 1)
        texto = texto.replace("```", "")
        texto = texto.strip()

    inicio = texto.find("{")
    fin = texto.rfind("}")

    if inicio != -1 and fin != -1 and fin > inicio:
        return texto[inicio:fin + 1]

    return texto


def limpiar_cuerpo_documento(doc):
    body = doc._body._element
    sect_pr = body.sectPr

    for elemento in list(body):
        body.remove(elemento)

    if sect_pr is not None:
        body.append(deepcopy(sect_pr))


def crear_documento_base(membrete_docx=None, membrete_path=None):
    if membrete_docx is not None:
        doc = Document(membrete_docx)
        limpiar_cuerpo_documento(doc)
    elif membrete_path and os.path.exists(membrete_path):
        doc = Document(membrete_path)
        limpiar_cuerpo_documento(doc)
    else:
        doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    return doc


def configurar_estilos(doc, fuente_base, tamanio_base, paleta, estilos_bd):
    styles = doc.styles
    color_texto = paleta.get("texto", COLOR_TEXTO_FALLBACK)

    try:
        normal = styles["Normal"]
        normal.font.name = fuente_base
        normal.font.size = Pt(tamanio_base)
        normal.font.color.rgb = RGBColor(*color_texto)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.1
    except KeyError:
        pass

    estilo_titulo = estilos_bd.get("titulo_1", {})

    for style_name in ["Heading 1", "Título 1"]:
        try:
            style = styles[style_name]
            style.font.name = estilo_titulo.get("tipografia") or fuente_base
            style.font.size = Pt(estilo_titulo.get("tamanio_letra") or 16)
            style.font.bold = estilo_titulo.get("negrita", True)
            style.font.italic = estilo_titulo.get("cursiva", False)
            style.font.color.rgb = RGBColor(*paleta.get("principal", COLOR_PRINCIPAL_FALLBACK))
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)
        except KeyError:
            continue

    estilo_titulo_2 = estilos_bd.get("titulo_2", {})

    for style_name in ["Heading 2", "Título 2", "Titulo 2", "Ttulo2"]:
        try:
            style = styles[style_name]
            style.font.name = estilo_titulo_2.get("tipografia") or fuente_base
            style.font.size = Pt(estilo_titulo_2.get("tamanio_letra") or 14)
            style.font.bold = estilo_titulo_2.get("negrita", True)
            style.font.italic = estilo_titulo_2.get("cursiva", False)
            style.font.color.rgb = RGBColor(*paleta.get("secundario", COLOR_PRINCIPAL_FALLBACK))
            style.paragraph_format.space_before = Pt(8)
            style.paragraph_format.space_after = Pt(4)
        except KeyError:
            continue


def agregar_membrete_simple(doc, nombre_empresa, nombre_cliente, fuente_base):
    section = doc.sections[0]

    header = section.header
    if not header.paragraphs:
        header.add_paragraph()

    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(nombre_empresa.upper())
    aplicar_fuente_run(run, fuente_base)
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*COLOR_PRINCIPAL_FALLBACK)

    p2 = header.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Documento de seguimiento de servicios especializados")
    aplicar_fuente_run(run2, fuente_base)
    run2.font.size = Pt(8)
    run2.font.color.rgb = RGBColor(90, 90, 90)

    footer = section.footer
    p_footer = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p_footer.text = ""
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_footer = p_footer.add_run(f"{nombre_empresa.upper()} | {nombre_cliente.upper()}")
    aplicar_fuente_run(run_footer, fuente_base)
    run_footer.font.size = Pt(8)
    run_footer.font.color.rgb = RGBColor(90, 90, 90)


def aplicar_bordes_tabla(tabla, color="000000"):
    tbl = tabla._tbl
    tbl_pr = tbl.tblPr

    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    borders = OxmlElement("w:tblBorders")

    for nombre in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{nombre}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), color)
        border.set(qn("w:space"), "0")
        borders.append(border)

    existente = tbl_pr.find(qn("w:tblBorders"))
    if existente is not None:
        tbl_pr.remove(existente)

    tbl_pr.append(borders)


def fijar_ancho_tabla(tabla, anchos_cm):
    for fila in tabla.rows:
        for idx, ancho in enumerate(anchos_cm):
            if idx >= len(fila.cells):
                continue
            celda = fila.cells[idx]
            celda.width = Cm(ancho)
            tc_pr = celda._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(Cm(ancho).twips)))
            tc_w.set(qn("w:type"), "dxa")


def colorear_celda(celda, color_rgb):
    tc_pr = celda._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), f"{color_rgb[0]:02X}{color_rgb[1]:02X}{color_rgb[2]:02X}")
    shading.set(qn("w:val"), "clear")
    tc_pr.append(shading)


def configurar_margenes_celda(celda, top=80, bottom=80, start=120, end=120):
    tc_pr = celda._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for nombre, valor in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        margen = tc_mar.find(qn(f"w:{nombre}"))
        if margen is None:
            margen = OxmlElement(f"w:{nombre}")
            tc_mar.append(margen)
        margen.set(qn("w:w"), str(valor))
        margen.set(qn("w:type"), "dxa")


def escribir_celda(
    celda,
    texto,
    bold=False,
    italic=False,
    color=COLOR_TEXTO_FALLBACK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    fuente_base=FUENTE_BASE_FALLBACK,
    tamanio_base=9,
):
    celda.text = ""
    celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    configurar_margenes_celda(celda)

    p = celda.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(texto))
    aplicar_fuente_run(run, fuente_base)
    run.font.size = Pt(tamanio_base)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)


def escribir_parrafo_en_celda(
    celda,
    texto,
    bold=False,
    italic=False,
    color=COLOR_TEXTO_FALLBACK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size=10,
    fuente_base=FUENTE_BASE_FALLBACK,
):
    p = celda.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(4)

    run = p.add_run(str(texto))
    aplicar_fuente_run(run, fuente_base)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)

    return p


def limpiar_celda(celda):
    celda.text = ""
    for parrafo in celda.paragraphs:
        parrafo.text = ""


def aplicar_alineacion(parrafo, alineacion):
    alineaciones = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    parrafo.alignment = alineaciones.get(
        (alineacion or "center").lower(),
        WD_ALIGN_PARAGRAPH.CENTER,
    )


def aplicar_heading_word(parrafo, nivel=1):
    estilos = {
        1: ["Heading 1", "Título 1"],
        2: ["Heading 2", "Título 2"],
        3: ["Heading 3", "Título 3"],
    }
    estilos.setdefault(nivel, []).extend([
        f"Heading {nivel}",
        f"Título {nivel}",
        f"Titulo {nivel}",
        f"Ttulo{nivel}",
    ])

    for nombre_estilo in estilos.get(nivel, []):
        try:
            parrafo.style = nombre_estilo
            aplicar_nivel_esquema(parrafo, nivel)
            return True
        except Exception:
            continue

    aplicar_nivel_esquema(parrafo, nivel)
    return False


def aplicar_nivel_esquema(parrafo, nivel):
    p_pr = parrafo._p.get_or_add_pPr()
    outline_lvl = p_pr.find(qn("w:outlineLvl"))

    if outline_lvl is None:
        outline_lvl = OxmlElement("w:outlineLvl")
        p_pr.append(outline_lvl)

    outline_lvl.set(qn("w:val"), str(max(nivel - 1, 0)))


def agregar_titulo(doc, titulo, estilo_titulo=None, color_principal=COLOR_PRINCIPAL_FALLBACK, fuente_base=FUENTE_BASE_FALLBACK):
    p = doc.add_paragraph()
    aplicar_heading_word(p, nivel=1)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)

    if estilo_titulo:
        aplicar_alineacion(p, estilo_titulo.get("alineacion") or "center")
        fuente = estilo_titulo.get("tipografia") or fuente_base
        tamanio = estilo_titulo.get("tamanio_letra") or 18
        negrita = estilo_titulo.get("negrita")
        cursiva = estilo_titulo.get("cursiva")
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fuente = fuente_base
        tamanio = 18
        negrita = True
        cursiva = False

    run = p.add_run(titulo.upper())
    aplicar_fuente_run(run, fuente)
    run.font.size = Pt(tamanio)
    run.font.bold = negrita
    run.font.italic = cursiva
    run.font.color.rgb = RGBColor(*color_principal)


def agregar_subtitulo(
    doc,
    titulo,
    estilo_titulo=None,
    color_secundario=COLOR_PRINCIPAL_FALLBACK,
    fuente_base=FUENTE_BASE_FALLBACK,
):
    p = doc.add_paragraph()
    aplicar_heading_word(p, nivel=2)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

    if estilo_titulo:
        aplicar_alineacion(p, estilo_titulo.get("alineacion") or "left")
        fuente = estilo_titulo.get("tipografia") or fuente_base
        tamanio = estilo_titulo.get("tamanio_letra") or 14
        negrita = estilo_titulo.get("negrita", True)
        cursiva = estilo_titulo.get("cursiva", False)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fuente = fuente_base
        tamanio = 14
        negrita = True
        cursiva = False

    run = p.add_run(titulo.upper())
    aplicar_fuente_run(run, fuente)
    run.font.size = Pt(tamanio)
    run.font.bold = negrita
    run.font.italic = cursiva
    run.font.color.rgb = RGBColor(*color_secundario)


def agregar_parrafo(
    doc,
    texto,
    alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
    fuente_base=FUENTE_BASE_FALLBACK,
    tamanio_base=TAMANIO_BASE_FALLBACK,
    color_texto=COLOR_TEXTO_FALLBACK,
):
    p = doc.add_paragraph()
    p.alignment = alineacion
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(texto)
    aplicar_fuente_run(run, fuente_base)
    run.font.size = Pt(tamanio_base)
    run.font.color.rgb = RGBColor(*color_texto)
    return p


def agregar_texto_estatico_seguimiento(
    doc,
    nombre_empresa,
    nombre_cliente,
    fuente_base,
    tamanio_base,
    color_texto,
):
    parrafos = [
        (
            f"En '{nombre_empresa.upper()}' es grato ofrecer el presente servicio "
            "y agradecer la confianza depositada para su debida realizacion."
        ),
        (
            f"Estimado '{nombre_cliente.upper()}' se busca garantizar la eficacia "
            "y calidad en los proyectos otorgados y alcance de los objetivos propuestos."
        ),
        (
            "Para cumplir y otorgar un resultado satisfactorio se supervisara y evaluara "
            "los avances en el servicio."
        ),
        (
            "Se asignaron temas que reforzaran la prestacion del servicio, esto con base "
            "a las necesidades de la empresa. Por ello, mediante esta tabla se mencionan "
            "dichos temas que seran abordados en el bimestre:"
        ),
    ]

    for texto in parrafos:
        agregar_parrafo(
            doc,
            texto,
            alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            color_texto=color_texto,
        )


def agregar_fecha_calendario(
    doc,
    fecha_documento,
    fuente_base,
    tamanio_base,
    color_texto,
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(26)

    run = p.add_run(f"Merida, Yucatan a {fecha_documento}")
    aplicar_fuente_run(run, fuente_base)
    run.font.size = Pt(tamanio_base)
    run.font.color.rgb = RGBColor(*color_texto)


def agregar_texto_estatico_calendario(
    doc,
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    fuente_base,
    tamanio_base,
    color_texto,
):
    agregar_parrafo(
        doc,
        f"Estimado {nombre_cliente.upper()}",
        alineacion=WD_ALIGN_PARAGRAPH.LEFT,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )

    agregar_parrafo(
        doc,
        (
            f"En {nombre_empresa.upper()} agradecemos la confianza que nos brinda "
            "para presentarle la propuesta de fechas para la realizacion del servicio "
            f"{nombre_servicio.upper()} que sera aplicado a la empresa."
        ),
        alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )


def agregar_tabla_calendario(
    doc,
    nombre_servicio,
    fecha_inicio,
    fecha_fin,
    paleta,
    fuente_base,
    tamanio_base,
):
    doc.add_paragraph()

    tabla = doc.add_table(rows=2, cols=4)
    tabla.autofit = False
    aplicar_bordes_tabla(tabla, color="000000")
    fijar_ancho_tabla(tabla, [1.0, 6.4, 3.5, 3.5])

    color_principal = paleta.get("principal", COLOR_PRINCIPAL_FALLBACK)
    encabezados = ["#", "Nombre del servicio", "Inicio", "Fin"]

    for idx, encabezado in enumerate(encabezados):
        celda = tabla.rows[0].cells[idx]
        escribir_celda(
            celda,
            encabezado,
            bold=True,
            color=(255, 255, 255),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
        )
        colorear_celda(celda, color_principal)

    valores = [
        "1",
        nombre_servicio.upper(),
        formatear_fecha_corta(fecha_inicio),
        formatear_fecha_corta(fecha_fin),
    ]

    for idx, valor in enumerate(valores):
        escribir_celda(
            tabla.rows[1].cells[idx],
            valor,
            align=WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.JUSTIFY,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
        )

    doc.add_paragraph()


def agregar_cierre_calendario(
    doc,
    fuente_base,
    tamanio_base,
    color_texto,
):
    parrafos = [
        (
            "Esperamos que la anterior propuesta de fechas para la realizacion del "
            "servicio se ajuste a sus tiempos y necesidades."
        ),
        (
            "Quedamos a sus ordenes por si tiene alguna duda que surja de la presente "
            "propuesta o cualquier comentario relacionado con el servicio."
        ),
        "Reciba un saludo afectuoso",
    ]

    for texto in parrafos:
        agregar_parrafo(
            doc,
            texto,
            alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            color_texto=color_texto,
        )


def agregar_firma_calendario(
    doc,
    nombre_cliente,
    fuente_base,
    tamanio_base,
    color_texto,
):
    doc.add_paragraph()

    p_atentamente = doc.add_paragraph()
    p_atentamente.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_atentamente.paragraph_format.space_after = Pt(20)
    run_atentamente = p_atentamente.add_run("ATENTAMENTE")
    aplicar_fuente_run(run_atentamente, fuente_base)
    run_atentamente.font.size = Pt(tamanio_base)
    run_atentamente.font.color.rgb = RGBColor(*color_texto)

    p_linea = doc.add_paragraph()
    p_linea.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_linea.paragraph_format.space_after = Pt(4)
    run_linea = p_linea.add_run("______________________________")
    aplicar_fuente_run(run_linea, fuente_base)
    run_linea.font.size = Pt(tamanio_base)
    run_linea.font.color.rgb = RGBColor(*color_texto)

    p_empresa = doc.add_paragraph()
    p_empresa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_empresa = p_empresa.add_run(nombre_cliente.upper())
    aplicar_fuente_run(run_empresa, fuente_base)
    run_empresa.font.size = Pt(tamanio_base)
    run_empresa.font.bold = True
    run_empresa.font.color.rgb = RGBColor(*color_texto)


def agregar_texto_estatico_acuse(
    doc,
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    periodo,
    fuente_base,
    tamanio_base,
    color_texto,
):
    agregar_parrafo(
        doc,
        "A quien corresponda:",
        alineacion=WD_ALIGN_PARAGRAPH.LEFT,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )
    doc.add_paragraph()

    agregar_parrafo(
        doc,
        f"Por medio de la presente hago constar que se le entrega a {nombre_cliente.upper()}",
        alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )

    agregar_parrafo(
        doc,
        (
            f"Manual completo de \"{nombre_servicio.upper()}\" elaborado por la empresa "
            f"{nombre_empresa.upper()}. en el bimestre de {periodo.upper()}."
        ),
        alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )


def agregar_firmas_acuse(
    doc,
    fuente_base,
    tamanio_base,
    color_texto,
):
    for _ in range(4):
        doc.add_paragraph()

    campos = [
        "NOMBRE DE QUIEN RECIBE:",
        "FIRMA DE QUIEN RECIBE:",
        "FECHA:",
    ]

    for campo in campos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(14)

        run_label = p.add_run(campo + " ")
        aplicar_fuente_run(run_label, fuente_base)
        run_label.font.size = Pt(tamanio_base)
        run_label.font.bold = True
        run_label.font.color.rgb = RGBColor(*color_texto)

        run_linea = p.add_run("______________________________")
        aplicar_fuente_run(run_linea, fuente_base)
        run_linea.font.size = Pt(tamanio_base)
        run_linea.font.color.rgb = RGBColor(*color_texto)


def agregar_portada_entregable(
    doc,
    nombre_servicio,
    fuente_base,
    tamanio_base,
    color_texto,
    estilo_titulo=None,
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(430)
    p.paragraph_format.space_after = Pt(0)

    fuente = fuente_base
    tamanio = max(tamanio_base + 8, 20)
    negrita = True

    if estilo_titulo:
        fuente = estilo_titulo.get("tipografia") or fuente
        tamanio = estilo_titulo.get("tamanio_letra") or tamanio
        negrita = estilo_titulo.get("negrita", True)

    run = p.add_run(nombre_servicio.upper())
    aplicar_fuente_run(run, fuente)
    run.font.size = Pt(tamanio)
    run.font.bold = negrita
    run.font.color.rgb = RGBColor(*color_texto)


def agregar_segunda_portada_entregable(
    doc,
    nombre_empresa,
    nombre_cliente,
    periodo,
    fuente_base,
    tamanio_base,
    color_texto,
    estilo_titulo=None,
):
    doc.add_page_break()
    for _ in range(27):
        p_espacio = doc.add_paragraph()
        p_espacio.paragraph_format.space_after = Pt(0)

    fuente = fuente_base
    tamanio = max(tamanio_base + 8, 20)

    if estilo_titulo:
        fuente = estilo_titulo.get("tipografia") or fuente
        tamanio = estilo_titulo.get("tamanio_letra") or tamanio

    lineas = [
        f"Dirigido a: {nombre_cliente.upper()}",
        f"Por: {nombre_empresa.upper()}",
        periodo.upper().replace(" - ", "-"),
    ]

    for idx, texto in enumerate(lineas):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0 if idx == 0 else 12)
        p.paragraph_format.space_after = Pt(0)

        run = p.add_run(texto)
        aplicar_fuente_run(run, fuente)
        run.font.size = Pt(tamanio)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*color_texto)


def habilitar_actualizacion_campos(doc):
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))

    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)

    update_fields.set(qn("w:val"), "true")


def agregar_campo_indice_word(parrafo):
    run_begin = parrafo.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fld_begin)

    run_instr = parrafo.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
    run_instr._r.append(instr_text)

    run_separate = parrafo.add_run()
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    run_separate._r.append(fld_separate)

    run_placeholder = parrafo.add_run(
        "Actualiza el indice en Word para mostrar los titulos del documento."
    )
    run_placeholder.font.italic = True
    run_placeholder.font.size = Pt(10)

    run_end = parrafo.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def agregar_indice_entregable(
    doc,
    fuente_base,
    tamanio_base,
    color_principal,
    color_texto,
    estilo_titulo=None,
):
    doc.add_page_break()

    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.paragraph_format.space_before = Pt(90)
    p_titulo.paragraph_format.space_after = Pt(24)

    fuente = fuente_base
    tamanio = max(tamanio_base + 6, 18)
    negrita = True
    cursiva = False

    if estilo_titulo:
        fuente = estilo_titulo.get("tipografia") or fuente
        tamanio = estilo_titulo.get("tamanio_letra") or tamanio
        negrita = estilo_titulo.get("negrita", True)
        cursiva = estilo_titulo.get("cursiva", False)

    run_titulo = p_titulo.add_run("CONTENIDO")
    aplicar_fuente_run(run_titulo, fuente)
    run_titulo.font.size = Pt(tamanio)
    run_titulo.font.bold = negrita
    run_titulo.font.italic = cursiva
    run_titulo.font.color.rgb = RGBColor(*color_principal)

    p_indice = doc.add_paragraph()
    p_indice.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_indice.paragraph_format.space_after = Pt(0)
    agregar_campo_indice_word(p_indice)

    for run in p_indice.runs:
        aplicar_fuente_run(run, fuente_base)
        if run.font.size is None:
            run.font.size = Pt(tamanio_base)
        if run.font.color.rgb is None:
            run.font.color.rgb = RGBColor(*color_texto)


def agregar_introduccion_entregable(
    doc,
    introduccion,
    fuente_base,
    tamanio_base,
    color_principal,
    color_texto,
    estilo_titulo=None,
):
    doc.add_page_break()

    agregar_titulo(
        doc,
        "Introduccion",
        estilo_titulo=estilo_titulo,
        color_principal=color_principal,
        fuente_base=fuente_base,
    )
    agregar_parrafo(
        doc,
        introduccion,
        alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )


def agregar_objetivos_entregable(
    doc,
    objetivos,
    fuente_base,
    tamanio_base,
    color_principal,
    color_secundario,
    color_texto,
    estilo_titulo_1=None,
    estilo_titulo_2=None,
):
    doc.add_page_break()

    agregar_titulo(
        doc,
        "Objetivo general",
        estilo_titulo=estilo_titulo_1,
        color_principal=color_principal,
        fuente_base=fuente_base,
    )
    agregar_parrafo(
        doc,
        objetivos["objetivo_general"],
        alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=color_texto,
    )
    agregar_subtitulo(
        doc,
        "Objetivos especificos",
        estilo_titulo=estilo_titulo_2,
        color_secundario=color_secundario,
        fuente_base=fuente_base,
    )

    for objetivo in objetivos["objetivos_especificos"]:
        agregar_parrafo(
            doc,
            f"♦\t{objetivo}",
            alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            color_texto=color_texto,
        )


def agregar_desarrollo_conceptos_entregable(
    doc,
    nombre_servicio,
    desarrollos,
    fuente_base,
    tamanio_base,
    color_principal,
    color_secundario,
    color_texto,
    estilo_titulo_1=None,
    estilo_titulo_2=None,
):
    doc.add_page_break()

    agregar_titulo(
        doc,
        nombre_servicio,
        estilo_titulo=estilo_titulo_1,
        color_principal=color_principal,
        fuente_base=fuente_base,
    )

    for idx, desarrollo in enumerate(desarrollos):
        if idx > 0:
            doc.add_page_break()

        agregar_subtitulo(
            doc,
            desarrollo["concepto"],
            estilo_titulo=estilo_titulo_2,
            color_secundario=color_secundario,
            fuente_base=fuente_base,
        )

        parrafos = [
            parrafo.strip()
            for parrafo in desarrollo["texto"].splitlines()
            if parrafo.strip()
        ]

        if not parrafos:
            parrafos = [desarrollo["texto"]]

        for parrafo in parrafos:
            agregar_parrafo(
                doc,
                parrafo,
                alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
                fuente_base=fuente_base,
                tamanio_base=tamanio_base,
                color_texto=color_texto,
            )


def agregar_referencias_entregable(
    doc,
    referencias,
    fuente_base,
    tamanio_base,
    color_principal,
    color_texto,
    estilo_titulo=None,
):
    doc.add_page_break()

    agregar_titulo(
        doc,
        "Referencias",
        estilo_titulo=estilo_titulo,
        color_principal=color_principal,
        fuente_base=fuente_base,
    )

    if not referencias:
        referencias = [
            "No se generaron referencias verificables para los conceptos desarrollados."
        ]

    for referencia in referencias:
        p = agregar_parrafo(
            doc,
            referencia,
            alineacion=WD_ALIGN_PARAGRAPH.LEFT,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            color_texto=color_texto,
        )
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_after = Pt(8)


def agregar_tabla_metodologia(
    doc,
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    periodo,
    metodologia,
    paleta,
    fuente_base,
    tamanio_base,
):
    tabla = doc.add_table(rows=5, cols=2)
    tabla.autofit = False
    aplicar_bordes_tabla(tabla, color="000000")
    fijar_ancho_tabla(tabla, [6.2, 9.3])

    color_principal = paleta.get("principal", COLOR_PRINCIPAL_FALLBACK)

    escribir_celda(
        tabla.rows[0].cells[0],
        "Nombre de empresa:",
        bold=True,
        italic=True,
        color=(255, 255, 255),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )
    colorear_celda(tabla.rows[0].cells[0], color_principal)

    escribir_celda(
        tabla.rows[0].cells[1],
        nombre_empresa.upper(),
        color=(255, 255, 255),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )
    colorear_celda(tabla.rows[0].cells[1], color_principal)

    escribir_celda(
        tabla.rows[1].cells[0],
        "Nombre de cliente:",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )
    escribir_celda(
        tabla.rows[1].cells[1],
        nombre_cliente.upper(),
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )

    celda_servicio = tabla.rows[2].cells[0].merge(tabla.rows[2].cells[1])
    escribir_celda(
        celda_servicio,
        nombre_servicio.upper(),
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )

    escribir_celda(
        tabla.rows[3].cells[0],
        "Metodologia:",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )

    celda_metodologia = tabla.rows[3].cells[1]
    limpiar_celda(celda_metodologia)
    configurar_margenes_celda(celda_metodologia, top=80, bottom=80, start=160, end=120)

    escribir_parrafo_en_celda(
        celda_metodologia,
        f"- {metodologia['tema']}",
        bold=True,
        color=paleta.get("texto", COLOR_TEXTO_FALLBACK),
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        size=tamanio_base,
        fuente_base=fuente_base,
    )

    for subtema in metodologia["subtemas"]:
        escribir_parrafo_en_celda(
            celda_metodologia,
            f"- {subtema}",
            color=paleta.get("texto", COLOR_TEXTO_FALLBACK),
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            size=tamanio_base,
            fuente_base=fuente_base,
        )

    escribir_celda(
        tabla.rows[4].cells[0],
        "Bimestre",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )
    escribir_celda(
        tabla.rows[4].cells[1],
        periodo.upper() + ".",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )


def crear_word_seguimiento(
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    periodo,
    metodologia,
    fuente_base,
    tamanio_base,
    membrete_docx=None,
    membrete_path=None,
    estilos_bd=None,
    paleta=None,
):
    paleta = paleta or {
        "principal": COLOR_PRINCIPAL_FALLBACK,
        "texto": COLOR_TEXTO_FALLBACK,
    }
    estilos_bd = estilos_bd or {}

    doc = crear_documento_base(membrete_docx, membrete_path)
    configurar_estilos(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        paleta=paleta,
        estilos_bd=estilos_bd,
    )

    if membrete_docx is None and not membrete_path:
        agregar_membrete_simple(doc, nombre_empresa, nombre_cliente, fuente_base)

    agregar_titulo(
        doc,
        "Seguimiento bimestral",
        estilo_titulo=estilos_bd.get("titulo_1"),
        color_principal=paleta["principal"],
        fuente_base=fuente_base,
    )

    agregar_texto_estatico_seguimiento(
        doc,
        nombre_empresa,
        nombre_cliente,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )
    agregar_tabla_metodologia(
        doc,
        nombre_empresa=nombre_empresa,
        nombre_cliente=nombre_cliente,
        nombre_servicio=nombre_servicio,
        periodo=periodo,
        metodologia=metodologia,
        paleta=paleta,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def crear_word_calendario(
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    fecha_documento,
    fecha_inicio,
    fecha_fin,
    fuente_base,
    tamanio_base,
    membrete_path=None,
    estilos_bd=None,
    paleta=None,
):
    paleta = paleta or {
        "principal": COLOR_PRINCIPAL_FALLBACK,
        "texto": COLOR_TEXTO_FALLBACK,
    }
    estilos_bd = estilos_bd or {}

    doc = crear_documento_base(membrete_path=membrete_path)
    configurar_estilos(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        paleta=paleta,
        estilos_bd=estilos_bd,
    )

    if not membrete_path:
        agregar_membrete_simple(doc, nombre_empresa, nombre_cliente, fuente_base)

    agregar_fecha_calendario(
        doc,
        fecha_documento=fecha_documento,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )
    agregar_titulo(
        doc,
        "Calendario de servicio",
        estilo_titulo=estilos_bd.get("titulo_1"),
        color_principal=paleta["principal"],
        fuente_base=fuente_base,
    )
    agregar_texto_estatico_calendario(
        doc,
        nombre_empresa=nombre_empresa,
        nombre_cliente=nombre_cliente,
        nombre_servicio=nombre_servicio,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )
    agregar_tabla_calendario(
        doc,
        nombre_servicio=nombre_servicio,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        paleta=paleta,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
    )
    agregar_cierre_calendario(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )
    agregar_firma_calendario(
        doc,
        nombre_cliente=nombre_cliente,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def crear_word_acuse(
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    periodo,
    fuente_base,
    tamanio_base,
    membrete_path=None,
    estilos_bd=None,
    paleta=None,
):
    paleta = paleta or {
        "principal": COLOR_PRINCIPAL_FALLBACK,
        "texto": COLOR_TEXTO_FALLBACK,
    }
    estilos_bd = estilos_bd or {}

    doc = crear_documento_base(membrete_path=membrete_path)
    configurar_estilos(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        paleta=paleta,
        estilos_bd=estilos_bd,
    )

    if not membrete_path:
        agregar_membrete_simple(doc, nombre_empresa, nombre_cliente, fuente_base)

    agregar_titulo(
        doc,
        "Acuse de recibo",
        estilo_titulo=estilos_bd.get("titulo_1"),
        color_principal=paleta["principal"],
        fuente_base=fuente_base,
    )
    agregar_texto_estatico_acuse(
        doc,
        nombre_empresa=nombre_empresa,
        nombre_cliente=nombre_cliente,
        nombre_servicio=nombre_servicio,
        periodo=periodo,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )
    agregar_firmas_acuse(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
    )

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def crear_word_entregable(
    nombre_empresa,
    nombre_cliente,
    nombre_servicio,
    periodo,
    introduccion,
    objetivos,
    desarrollo_conceptos,
    referencias,
    fuente_base,
    tamanio_base,
    membrete_path=None,
    estilos_bd=None,
    paleta=None,
):
    paleta = paleta or {
        "principal": COLOR_PRINCIPAL_FALLBACK,
        "texto": COLOR_TEXTO_FALLBACK,
    }
    estilos_bd = estilos_bd or {}

    doc = crear_documento_base(membrete_path=membrete_path)
    configurar_estilos(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        paleta=paleta,
        estilos_bd=estilos_bd,
    )

    if not membrete_path:
        agregar_membrete_simple(doc, nombre_empresa, nombre_cliente, fuente_base)

    habilitar_actualizacion_campos(doc)

    agregar_portada_entregable(
        doc,
        nombre_servicio=nombre_servicio,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
        estilo_titulo=estilos_bd.get("titulo_1"),
    )
    agregar_segunda_portada_entregable(
        doc,
        nombre_empresa=nombre_empresa,
        nombre_cliente=nombre_cliente,
        periodo=periodo,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_texto=paleta["texto"],
        estilo_titulo=estilos_bd.get("titulo_1"),
    )
    agregar_indice_entregable(
        doc,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_principal=paleta["principal"],
        color_texto=paleta["texto"],
        estilo_titulo=estilos_bd.get("titulo_1"),
    )
    agregar_introduccion_entregable(
        doc,
        introduccion=introduccion,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_principal=paleta["principal"],
        color_texto=paleta["texto"],
        estilo_titulo=estilos_bd.get("titulo_1"),
    )
    agregar_objetivos_entregable(
        doc,
        objetivos=objetivos,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_principal=paleta["principal"],
        color_secundario=paleta.get("secundario", paleta["principal"]),
        color_texto=paleta["texto"],
        estilo_titulo_1=estilos_bd.get("titulo_1"),
        estilo_titulo_2=estilos_bd.get("titulo_2"),
    )
    agregar_desarrollo_conceptos_entregable(
        doc,
        nombre_servicio=nombre_servicio,
        desarrollos=desarrollo_conceptos,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_principal=paleta["principal"],
        color_secundario=paleta.get("secundario", paleta["principal"]),
        color_texto=paleta["texto"],
        estilo_titulo_1=estilos_bd.get("titulo_1"),
        estilo_titulo_2=estilos_bd.get("titulo_2"),
    )
    agregar_referencias_entregable(
        doc,
        referencias=referencias,
        fuente_base=fuente_base,
        tamanio_base=tamanio_base,
        color_principal=paleta["principal"],
        color_texto=paleta["texto"],
        estilo_titulo=estilos_bd.get("titulo_1"),
    )

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def mostrar_modulo_triple_a():
    st.title("Generador de seguimiento AAA")

    st.write(
        "Genera un documento Word de seguimiento con membrete, textos estaticos, "
        "tablas codificadas en Python y textos opcionales generados con IA."
    )

    empresas = obtener_empresas_triple_a()

    if not empresas:
        st.error("No hay empresas registradas en la BD.")
        return

    empresas_por_nombre = {
        empresa["nombre"]: empresa
        for empresa in empresas
    }

    with st.form("form_triple_a"):
        col1, col2 = st.columns(2)

        with col1:
            empresa_seleccionada = st.selectbox(
                "Empresa prestadora",
                list(empresas_por_nombre.keys()),
            )
            nombre_servicio = st.text_input(
                "Nombre del servicio",
                value="",
            )
            archivos_xml = st.file_uploader(
                "XML para calcular fecha del documento",
                type=["xml"],
                accept_multiple_files=True,
                help=(
                    "Se leera la etiqueta Fecha de cada XML, se ordenaran las fechas "
                    "y se restaran 4 dias habiles a la fecha mas antigua."
                ),
            )

        with col2:
            nombre_cliente = st.text_input(
                "Cliente",
                value="",
            )

        st.markdown("**Periodo**")

        col_periodo_inicio, col_periodo_fin, col_periodo_anio = st.columns(3)

        with col_periodo_inicio:
            mes_inicio_periodo = st.selectbox(
                "Mes inicio",
                MESES_SELECT,
                index=None,
                placeholder="Selecciona mes",
                format_func=lambda item: item[1],
                key="mes_inicio_periodo_triple_a",
            )

        with col_periodo_fin:
            mes_fin_periodo = st.selectbox(
                "Mes fin",
                MESES_SELECT,
                index=None,
                placeholder="Selecciona mes",
                format_func=lambda item: item[1],
                key="mes_fin_periodo_triple_a",
            )

        with col_periodo_anio:
            anio_periodo = st.selectbox(
                "Año",
                list(range(date.today().year - 5, date.today().year + 6)),
                index=None,
                placeholder="Selecciona año",
                key="anio_periodo_triple_a",
            )

        generar = st.form_submit_button("Generar documento")

    if not generar:
        return

    empresa_data = empresas_por_nombre[empresa_seleccionada]
    nombre_empresa = empresa_data["razon_social"]
    estilos_bd = obtener_estilos_word_empresa(empresa_data["id"])
    membrete_path = empresa_data["membrete_path"]
    fuente_base = empresa_data["tipografia_base"] or FUENTE_BASE_FALLBACK
    tamanio_base = int(empresa_data["tamanio_base"] or TAMANIO_BASE_FALLBACK)
    paleta = {
        "principal": convertir_color_bd_a_rgb(empresa_data["color_primario"]),
        "secundario": convertir_color_bd_a_rgb(empresa_data["color_secundario"]),
        "acento": convertir_color_bd_a_rgb(empresa_data["color_acento"]),
        "texto": COLOR_TEXTO_FALLBACK,
    }

    if not nombre_cliente.strip():
        st.warning("Escribe el cliente.")
        return

    if not nombre_servicio.strip():
        st.warning("Escribe el nombre del servicio.")
        return

    if not archivos_xml:
        st.warning("Sube al menos un XML para calcular la fecha del documento.")
        return

    if mes_inicio_periodo is None:
        st.warning("Selecciona el mes de inicio del periodo.")
        return

    if mes_fin_periodo is None:
        st.warning("Selecciona el mes fin del periodo.")
        return

    if anio_periodo is None:
        st.warning("Selecciona el anio del periodo.")
        return

    if mes_inicio_periodo[0] > mes_fin_periodo[0]:
        st.warning("El mes inicio no puede ser posterior al mes fin.")
        return

    try:
        fecha_documento, fechas_xml = calcular_fecha_documento_desde_xml(archivos_xml)
    except ValueError as e:
        st.error(str(e))
        return

    st.info(
        "Fechas XML detectadas: "
        + ", ".join(fecha.strftime("%d/%m/%Y") for fecha in fechas_xml)
        + f". Fecha del documento: {fecha_documento}."
    )

    periodo = (
        f"{MESES_ES[mes_inicio_periodo[0]]} - "
        f"{MESES_ES[mes_fin_periodo[0]]} {anio_periodo}"
    )

    try:
        with st.spinner("Generando metodologia, introduccion, objetivos, desarrollo de conceptos y referencias con IA..."):
            metodologia = generar_metodologia_ia(
                nombre_servicio=nombre_servicio,
                nombre_cliente=nombre_cliente,
            )
            conceptos_metodologia = obtener_conceptos_metodologia(metodologia)
            introduccion_entregable = generar_introduccion_entregable_ia(
                nombre_servicio=nombre_servicio,
                tema_metodologia=metodologia["tema"],
                nombre_cliente=nombre_cliente,
            )
            objetivos_entregable = generar_objetivos_entregable_ia(
                nombre_servicio=nombre_servicio,
                tema_metodologia=metodologia["tema"],
                nombre_cliente=nombre_cliente,
            )
            desarrollo_conceptos_entregable = generar_desarrollo_conceptos_entregable_ia(
                nombre_servicio=nombre_servicio,
                conceptos=conceptos_metodologia,
                nombre_cliente=nombre_cliente,
            )
            referencias_entregable = generar_referencias_entregable_ia(
                nombre_servicio=nombre_servicio,
                desarrollo_conceptos=desarrollo_conceptos_entregable,
            )

        word = crear_word_seguimiento(
            nombre_empresa=nombre_empresa,
            nombre_cliente=nombre_cliente,
            nombre_servicio=nombre_servicio,
            periodo=periodo,
            metodologia=metodologia,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            membrete_path=membrete_path,
            estilos_bd=estilos_bd,
            paleta=paleta,
        )
        calendario = crear_word_calendario(
            nombre_empresa=nombre_empresa,
            nombre_cliente=nombre_cliente,
            nombre_servicio=nombre_servicio,
            fecha_documento=fecha_documento,
            fecha_inicio=fechas_xml[0],
            fecha_fin=fechas_xml[-1],
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            membrete_path=membrete_path,
            estilos_bd=estilos_bd,
            paleta=paleta,
        )
        acuse = crear_word_acuse(
            nombre_empresa=nombre_empresa,
            nombre_cliente=nombre_cliente,
            nombre_servicio=nombre_servicio,
            periodo=periodo,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            membrete_path=membrete_path,
            estilos_bd=estilos_bd,
            paleta=paleta,
        )
        entregable = crear_word_entregable(
            nombre_empresa=nombre_empresa,
            nombre_cliente=nombre_cliente,
            nombre_servicio=nombre_servicio,
            periodo=periodo,
            introduccion=introduccion_entregable,
            objetivos=objetivos_entregable,
            desarrollo_conceptos=desarrollo_conceptos_entregable,
            referencias=referencias_entregable,
            fuente_base=fuente_base,
            tamanio_base=tamanio_base,
            membrete_path=membrete_path,
            estilos_bd=estilos_bd,
            paleta=paleta,
        )

        st.success("Documentos generados correctamente.")
        col_descarga_1, col_descarga_2, col_descarga_3, col_descarga_4 = st.columns(4)

        with col_descarga_1:
            st.download_button(
                label="Descargar seguimiento Word",
                data=word,
                file_name="1_seguimiento.docx",
                mime=MIME_DOCX,
                key="descargar_triple_a_seguimiento",
                on_click="ignore",
            )

        with col_descarga_2:
            st.download_button(
                label="Descargar calendario Word",
                data=calendario,
                file_name="2_calendario.docx",
                mime=MIME_DOCX,
                key="descargar_triple_a_calendario",
                on_click="ignore",
            )

        with col_descarga_3:
            st.download_button(
                label="Descargar entregable Word",
                data=entregable,
                file_name="3_entregable.docx",
                mime=MIME_DOCX,
                key="descargar_triple_a_entregable",
                on_click="ignore",
            )

        with col_descarga_4:
            st.download_button(
                label="Descargar acuse Word",
                data=acuse,
                file_name="4_acuse.docx",
                mime=MIME_DOCX,
                key="descargar_triple_a_acuse",
                on_click="ignore",
            )

    except ValueError as e:
        st.error(str(e))
    except Exception as e:
        st.error(f"Ocurrio un error al generar el documento: {e}")


if __name__ == "__main__":
    mostrar_modulo_triple_a()
